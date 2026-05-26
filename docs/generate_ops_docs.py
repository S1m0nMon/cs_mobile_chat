#!/usr/bin/env python3
"""
generate_ops_docs.py — HireVisa Mobile CS Chat
운영팀 업무 요청용 DOCX + 작성 양식 XLSX 생성

Run: python3 docs/generate_ops_docs.py
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ══════════════════════════════════════════════════════════════════
# 공통 색상 팔레트
# ══════════════════════════════════════════════════════════════════
NAVY      = "1F4E79"
NAVY_LT   = "D6E4F0"
AMBER     = "F59E0B"
AMBER_LT  = "FEF3C7"
EMERALD   = "059669"
EMERALD_LT= "D1FAE5"
PURPLE    = "7C3AED"
PURPLE_LT = "EDE9FE"
BLUE_LT   = "DBEAFE"
RED_LT    = "FEE2E2"
RED       = "DC2626"
GRAY_LT   = "F3F4F6"
GRAY_MID  = "E5E7EB"
GRAY      = "6B7280"
WHITE     = "FFFFFF"
YELLOW    = "FFF9C4"   # 입력 필요 셀
TEAL_LT   = "CCFBF1"
ORANGE_LT = "FFF7ED"
ORANGE    = "EA580C"

# ══════════════════════════════════════════════════════════════════
# XLSX 헬퍼
# ══════════════════════════════════════════════════════════════════
def xfill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def xfont(bold=False, color="1F2937", size=10, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic, name="Malgun Gothic")

def xborder(color="D1D5DB"):
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def xcenter():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def xleft():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def xh(ws, row, col, text, bg=NAVY, fg=WHITE, bold=True, size=10, colspan=1):
    """Header cell"""
    c = ws.cell(row=row, column=col, value=text)
    c.fill = xfill(bg)
    c.font = xfont(bold=bold, color=fg, size=size)
    c.alignment = xcenter()
    c.border = xborder()
    if colspan > 1:
        ws.merge_cells(start_row=row, start_column=col,
                       end_row=row, end_column=col + colspan - 1)
    return c

def xc(ws, row, col, text, bg=WHITE, fg="1F2937", bold=False,
       align="left", size=10, italic=False):
    """Data cell"""
    c = ws.cell(row=row, column=col, value=text)
    c.fill = xfill(bg)
    c.font = xfont(bold=bold, color=fg, size=size, italic=italic)
    c.alignment = xleft() if align == "left" else xcenter()
    c.border = xborder()
    return c

def xinput(ws, row, col, placeholder="", bg=YELLOW):
    """Input cell (노란색 = 작성 필요)"""
    c = ws.cell(row=row, column=col, value=placeholder)
    c.fill = xfill(bg)
    c.font = xfont(color=GRAY, italic=True, size=10)
    c.alignment = xleft()
    c.border = xborder("F59E0B")
    return c

def sheet_banner(ws, title, subtitle=""):
    ws.row_dimensions[1].height = 34
    val = title + (f"   |   {subtitle}" if subtitle else "")
    c = ws.cell(row=1, column=1, value=val)
    c.fill = xfill(NAVY)
    c.font = xfont(bold=True, color=WHITE, size=13)
    c.alignment = xcenter()
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=20)

def note_row(ws, row, text, ncols=10):
    c = ws.cell(row=row, column=1, value=text)
    c.fill = xfill(GRAY_LT)
    c.font = xfont(italic=True, color=GRAY, size=9)
    c.alignment = xleft()
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.row_dimensions[row].height = 20

def input_legend(ws, start_row, ncols=10):
    """노란색 셀 = 작성 필요 범례"""
    r = start_row
    ws.row_dimensions[r].height = 8
    r += 1
    c = ws.cell(row=r, column=1, value="")
    c.fill = xfill(YELLOW)
    c.border = xborder("F59E0B")
    c = ws.cell(row=r, column=2,
                value="← 노란색 셀: 운영팀 직접 작성 필요   |   흰색 셀: 참고용 (현재 시스템 기준)")
    c.font = xfont(italic=True, color=AMBER, size=9, bold=True)
    c.alignment = xleft()
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncols)


# ══════════════════════════════════════════════════════════════════
# XLSX Sheet 1: 안내 & 체크리스트
# ══════════════════════════════════════════════════════════════════
def build_guide(wb):
    ws = wb.create_sheet("📌 안내 및 체크리스트")
    sheet_banner(ws, "📌 운영팀 작성 안내 및 체크리스트",
                 "작성 전 반드시 읽어주세요")

    # 안내문
    ws.row_dimensions[2].height = 8
    intro_texts = [
        ("이 문서의 목적",
         "HireVisa Mobile CS 채팅 시스템의 운영 기준을 확정하기 위해 운영팀의 의견을 수렴합니다.\n"
         "각 시트에 노란색(🟡)으로 표시된 셀은 운영팀이 직접 작성해야 하는 항목입니다.\n"
         "흰색 셀은 현재 시스템에 임시로 설정된 값이며, 검토 후 수정 가능합니다."),
        ("채팅 시스템 개요",
         "학생이 채팅으로 진입하면 ① 신청번호 입력 또는 ② 채팅 내 간편신청을 통해 상담이 시작됩니다.\n"
         "오퍼레이터는 고객의 Seg(개통 유형 × 신분증 유형)를 확인하고, 서류 체크리스트를 통해 승인/보완 처리를 진행합니다.\n"
         "모든 서류가 승인되면 개통 완료 처리가 가능합니다."),
        ("Seg 구분 (4가지)",
         "① 여권 신규: 여권 소지 + 신규 번호 개통\n"
         "② 여권 번이: 여권 소지 + 기존 번호 이동 (MNP)\n"
         "③ 외등 신규: 외국인등록증 소지 + 신규 번호 개통\n"
         "④ 외등 번이: 외국인등록증 소지 + 기존 번호 이동 (MNP)\n"
         "각 Seg에 따라 필요 서류와 응대 시나리오가 달라집니다."),
        ("SIM 유형",
         "eSIM(디지털): 기기에 내장, 기기정보 캡처본 필요\n"
         "uSIM(실물): 물리 SIM 카드, 유심 사진 필요\n"
         "SIM 유형은 Seg와 독립적으로 결합됩니다. (예: 여권 신규 × eSIM / 여권 신규 × uSIM)"),
        ("자동 발송 메시지",
         "특정 상황(서류 전체 승인, 개통 완료 등)에서 오퍼레이터가 버튼을 누르지 않아도 자동으로 메시지가 발송됩니다.\n"
         "현재는 임시 문구가 설정되어 있으며, 운영팀이 공식 문구로 교체해야 합니다."),
    ]

    r = 3
    for title, body in intro_texts:
        ws.row_dimensions[r].height = 16
        c = ws.cell(row=r, column=1, value=title)
        c.fill = xfill(NAVY)
        c.font = xfont(bold=True, color=WHITE, size=10)
        c.alignment = xcenter()
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        r += 1

        c = ws.cell(row=r, column=1, value=body)
        c.fill = xfill(NAVY_LT)
        c.font = xfont(size=9)
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        lines = body.count("\n") + 1
        ws.row_dimensions[r].height = max(30, lines * 14)
        r += 1
        ws.row_dimensions[r].height = 6
        r += 1

    # 체크리스트 테이블
    ws.row_dimensions[r].height = 14
    r += 1
    ws.row_dimensions[r].height = 20
    xh(ws, r, 1, "✅ 작성 체크리스트", colspan=10)
    r += 1

    check_headers = ["#", "시트명", "작성 내용 요약", "담당", "완료 여부", "비고"]
    check_widths = [4, 24, 50, 12, 12, 20]
    for col, (h, w) in enumerate(zip(check_headers, check_widths), 1):
        xh(ws, r, col, h, bg=GRAY_LT, fg="1F2937", size=9)
        ws.column_dimensions[get_column_letter(col)].width = w
    r += 1

    checks = [
        ["1", "🗂 Seg별 응대 시나리오",
         "4개 Seg별 응대 흐름, 특이사항, 예상 처리 시간 정의", "", ""],
        ["2", "📄 필요 서류 정의",
         "Seg별 필수/선택 서류 확정 및 허용 기준 정의", "", ""],
        ["3", "💬 채팅 템플릿 메시지",
         "5개 단계별 공식 발송 문구 작성 (현재 임시 문구 검토 포함)", "", ""],
        ["4", "⚙️ 자동화 워크플로우",
         "자동 메시지 발송 시점 및 조건 확인·보완", "", ""],
        ["5", "📨 보완 요청 표준 문구",
         "항목별 보완 요청 시 표준 안내 문구 작성", "", ""],
        ["6", "❓ 퀵리플라이 응답 스크립트",
         "학생 버튼 클릭 시 오퍼레이터 표준 응답 스크립트 작성", "", ""],
    ]
    for row_data in checks:
        xc(ws, r, 1, row_data[0], bg=GRAY_LT, align="center", bold=True)
        xc(ws, r, 2, row_data[1], bg=GRAY_LT, bold=True)
        xc(ws, r, 3, row_data[2])
        xinput(ws, r, 4, "(담당자 이름)")
        xinput(ws, r, 5, "□ 미완료")
        xinput(ws, r, 6, "")
        ws.row_dimensions[r].height = 28
        r += 1

    ws.column_dimensions["A"].width = 4
    input_legend(ws, r + 1, ncols=10)


# ══════════════════════════════════════════════════════════════════
# XLSX Sheet 2: Seg별 응대 시나리오
# ══════════════════════════════════════════════════════════════════
def build_seg_scenario(wb):
    ws = wb.create_sheet("🗂 Seg별 응대 시나리오")
    sheet_banner(ws, "🗂 Seg별 응대 시나리오",
                 "4개 Seg × 응대 흐름·특이사항·주의사항 정의")

    headers = ["Seg", "대상 고객", "신분증 유형", "개통 유형",
               "응대 순서 (단계별 정리)", "특이사항 / 자주 발생 이슈",
               "예상 처리 시간", "오퍼레이터 주의사항"]
    widths  = [14,    20,           14,             14,
               50,                  40,
               16,                  40]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        xh(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    segs = [
        ("여권 신규", "한국 미거주 외국인 유학생\n(여권만 소지)",
         "여권 (Passport)", "신규 번호 개통 (S1)",
         BLUE_LT),
        ("여권 번이", "기존 한국 번호 보유 유학생\n(여권만 소지)",
         "여권 (Passport)", "번호 이동 MNP (S2)",
         AMBER_LT),
        ("외등 신규", "외국인등록증 소지 유학생\n(장기 거주자)",
         "외국인등록증 (ARC)", "신규 번호 개통 (S1)",
         EMERALD_LT),
        ("외등 번이", "외국인등록증 소지 유학생\n(장기 거주자·번호 이동)",
         "외국인등록증 (ARC)", "번호 이동 MNP (S2)",
         PURPLE_LT),
    ]

    for r_idx, (seg, target, id_type, act_type, bg) in enumerate(segs, 4):
        xc(ws, r_idx, 1, seg, bg=bg, bold=True, align="center")
        xc(ws, r_idx, 2, target, bg=bg)
        xc(ws, r_idx, 3, id_type, bg=bg, align="center")
        xc(ws, r_idx, 4, act_type, bg=bg, align="center")
        xinput(ws, r_idx, 5,
               "예시:\n1. 신청번호 입력 확인\n2. Seg 확인 (오퍼레이터)\n3. 서류 검토\n4. 보완 요청(필요시)\n5. 전체 승인 후 개통")
        xinput(ws, r_idx, 6, "예시: 여권 만료 케이스 처리 방법, 통신사 확인 불가 시 처리 등")
        xinput(ws, r_idx, 7, "예시: 평균 30분")
        xinput(ws, r_idx, 8, "예시: 여권 만료일 반드시 확인")
        ws.row_dimensions[r_idx].height = 90

    note_row(ws, 8,
             "💡 '응대 순서'는 오퍼레이터 업무 매뉴얼 기준으로 작성. "
             "시스템 자동화 부분(자동 메시지 발송 등)은 ⚙️ 자동화 워크플로우 시트 참고.", ncols=8)
    input_legend(ws, 10, ncols=8)


# ══════════════════════════════════════════════════════════════════
# XLSX Sheet 3: 필요 서류 정의
# ══════════════════════════════════════════════════════════════════
def build_docs_matrix(wb):
    ws = wb.create_sheet("📄 필요 서류 정의")
    sheet_banner(ws, "📄 Seg별 필요 서류 정의",
                 "현재 시스템 기준 (흰색) + 운영팀 확정 필요 (노란색)")

    headers = ["Seg", "서류명", "필수/선택", "제출 형태",
               "허용 포맷", "확인 기준", "보완 요청 기준 (요약)", "비고"]
    widths  = [14,    22,       10,          14,
               16,    30,       36,                             20]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        xh(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    # 현재 시스템 기준 서류 (white = 참고용)
    items = [
        # 여권 신규 (S1·여권)
        ("여권 신규", "여권 양면 평판 스캔본", "필수", "파일(이미지/PDF)",
         "JPG, PNG, PDF", "정보면+비자면 양면 포함, 정보 식별 가능, 유효기간 확인",
         "흐릿함/한쪽면/정보가림/만료", "", BLUE_LT),
        ("여권 신규", "기기정보 캡처본 (eSIM 전용)", "eSIM 필수 / uSIM 불필요", "파일(이미지)",
         "JPG, PNG", "IMEI 포함, 기기 모델명 확인 가능",
         "IMEI 미확인/모델불가/eSIM지원불가", "", BLUE_LT),
        ("여권 신규", "유심 사진 (uSIM 전용)", "uSIM 필수 / eSIM 불필요", "파일(이미지)",
         "JPG, PNG", "유심 앞면 선명, 유심 번호 확인 가능",
         "흐릿함/번호미확인/다른유심", "", BLUE_LT),

        # 여권 번이 (S2·여권)
        ("여권 번이", "여권 양면 평판 스캔본", "필수", "파일(이미지/PDF)",
         "JPG, PNG, PDF", "정보면+비자면 양면 포함, 정보 식별 가능, 유효기간 확인",
         "흐릿함/한쪽면/정보가림/만료", "", AMBER_LT),
        ("여권 번이", "기존 통신사 확인 문자 캡처", "필수", "파일(이미지)",
         "JPG, PNG", "114 발송 통신사 안내 문자, 번호 일치",
         "불명확/통신사불확인/다른번호", "", AMBER_LT),
        ("여권 번이", "기존 전화번호", "필수", "텍스트 입력",
         "010-XXXX-XXXX", "한국 휴대폰 번호 형식, 통신사 문자와 일치",
         "형식오류/번호불일치/명의불일치", "", AMBER_LT),
        ("여권 번이", "기기정보 캡처본 (eSIM 전용)", "eSIM 필수", "파일(이미지)",
         "JPG, PNG", "IMEI 포함, 기기 모델명 확인 가능",
         "IMEI 미확인/모델불가/eSIM지원불가", "", AMBER_LT),
        ("여권 번이", "유심 사진 (uSIM 전용)", "uSIM 필수", "파일(이미지)",
         "JPG, PNG", "유심 앞면 선명, 유심 번호 확인 가능",
         "흐릿함/번호미확인", "", AMBER_LT),

        # 외등 신규 (S1·외등)
        ("외등 신규", "외국인등록증 스캔본 (앞면)", "필수", "파일(이미지/PDF)",
         "JPG, PNG, PDF", "(운영팀 확정 필요)",
         "(운영팀 확정 필요)", "※ 현재 미구현 — 서류 기준 확정 필요", EMERALD_LT),
        ("외등 신규", "외국인등록증 스캔본 (뒷면)", "필수/선택?", "파일(이미지/PDF)",
         "JPG, PNG, PDF", "(운영팀 확정 필요)",
         "(운영팀 확정 필요)", "※ 현재 미구현", EMERALD_LT),
        ("외등 신규", "기기정보 캡처본 (eSIM 전용)", "eSIM 필수", "파일(이미지)",
         "JPG, PNG", "IMEI 포함, 기기 모델명 확인 가능",
         "IMEI 미확인/모델불가", "", EMERALD_LT),
        ("외등 신규", "유심 사진 (uSIM 전용)", "uSIM 필수", "파일(이미지)",
         "JPG, PNG", "유심 앞면 선명, 유심 번호 확인 가능",
         "흐릿함/번호미확인", "", EMERALD_LT),

        # 외등 번이 (S2·외등)
        ("외등 번이", "외국인등록증 스캔본 (앞면)", "필수", "파일(이미지/PDF)",
         "JPG, PNG, PDF", "(운영팀 확정 필요)",
         "(운영팀 확정 필요)", "※ 현재 미구현", PURPLE_LT),
        ("외등 번이", "기존 통신사 확인 문자 캡처", "필수", "파일(이미지)",
         "JPG, PNG", "114 발송 통신사 안내 문자, 번호 일치",
         "불명확/통신사불확인", "", PURPLE_LT),
        ("외등 번이", "기존 전화번호", "필수", "텍스트 입력",
         "010-XXXX-XXXX", "한국 휴대폰 번호 형식",
         "형식오류/번호불일치", "", PURPLE_LT),
        ("외등 번이", "기기정보 캡처본 (eSIM 전용)", "eSIM 필수", "파일(이미지)",
         "JPG, PNG", "IMEI 포함",
         "IMEI 미확인", "", PURPLE_LT),
        ("외등 번이", "유심 사진 (uSIM 전용)", "uSIM 필수", "파일(이미지)",
         "JPG, PNG", "유심 앞면 선명",
         "흐릿함/번호미확인", "", PURPLE_LT),
        # 선택 항목
        ("(선택 추가)", "사증발급확인서", "오퍼레이터 판단", "파일(이미지/PDF)",
         "JPG, PNG, PDF", "법무부 발급, 유효기간 확인",
         "불명확/만료/다른서류", "자사 미확보 시만 요청", GRAY_LT),
    ]

    for r_idx, row in enumerate(items, 4):
        seg, name, req, form, fmt, criteria, supp, note, bg = row
        xc(ws, r_idx, 1, seg,  bg=bg, bold=True, align="center")
        xc(ws, r_idx, 2, name, bg=bg, bold=True)
        xc(ws, r_idx, 3, req,  bg=bg, align="center")
        xc(ws, r_idx, 4, form, bg=bg, align="center")
        xc(ws, r_idx, 5, fmt,  bg=bg, align="center")
        # 확인 기준 — 외등은 입력 필요
        if "(운영팀 확정 필요)" in criteria:
            xinput(ws, r_idx, 6, criteria)
        else:
            xc(ws, r_idx, 6, criteria, bg=bg)
        # 보완 요청 기준 — 외등은 입력 필요
        if "(운영팀 확정 필요)" in supp:
            xinput(ws, r_idx, 7, supp)
        else:
            xc(ws, r_idx, 7, supp, bg=bg)
        xc(ws, r_idx, 8, note, bg=WHITE, fg=GRAY, italic=True)
        ws.row_dimensions[r_idx].height = 36

    note_row(ws, len(items) + 4 + 1,
             "💡 외등(외국인등록증) Seg는 현재 시스템 미구현 상태입니다. "
             "운영팀에서 서류 기준을 확정하면 개발팀이 시스템에 반영합니다.", ncols=8)
    input_legend(ws, len(items) + 4 + 3, ncols=8)


# ══════════════════════════════════════════════════════════════════
# XLSX Sheet 4: 채팅 템플릿 메시지
# ══════════════════════════════════════════════════════════════════
def build_templates(wb):
    ws = wb.create_sheet("💬 채팅 템플릿 메시지")
    sheet_banner(ws, "💬 채팅 템플릿 메시지 (단계별)",
                 "임시 문구(흰색) 검토 후 공식 문구(노란색)로 작성")

    headers = ["단계", "발송 시점 (자동/수동)", "발송 조건",
               "현재 임시 문구 (참고용)", "공식 문구 (운영팀 작성)", "변수", "비고"]
    widths  = [18,     20,                    32,
               48,                            48,                    16,    16]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        xh(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    templates = [
        ("채팅 시작\n인트로",
         "자동 발송 (300ms)\n앱번호 경로",
         "오퍼레이터가 Seg 최초 확인 시",
         "안녕하세요, {이름}님! 하이어비자 모바일 개통팀입니다.\n"
         "신청해 주셔서 감사합니다. 아래 서류 목록을 확인하시고 순서대로 제출해 주세요. "
         "궁금한 점이 있으시면 언제든지 말씀해 주세요. 🙏",
         "", "{이름}", "워크플로우 경로는 1800ms 딜레이", NAVY_LT),

        ("서류 검토 중\n(재제출 확인)",
         "자동 발송 (600ms)\n학생 재제출 후",
         "학생이 보완 요청 서류를 재제출하고 Status가 C로 전환될 때",
         "재제출해 주신 서류를 확인하겠습니다. 잠시만 기다려 주세요. 🔍",
         "", "없음", "여러 항목 재제출 시 마지막 항목 기준 1회 발송", BLUE_LT),

        ("개통 준비 완료\n(전체 승인)",
         "자동 발송 (400ms)\n마지막 항목 승인 후",
         "모든 서류 항목 승인 완료 → Status D 전환 시",
         "모든 서류 확인이 완료되었습니다! 곧 개통 처리를 진행해 드리겠습니다. 📱",
         "", "없음", "개통 완료 버튼 노출 전 발송", EMERALD_LT),

        ("개통 완료",
         "자동 발송 (즉시)\n오퍼레이터 클릭 후",
         "'개통 완료 처리' 버튼 클릭 시 (Status D 조건)",
         "{이름}님, 모바일 개통이 정상적으로 완료되었습니다! 🎉\n"
         "{개통유형} · {SIM유형} 방식으로 개통 처리되었습니다.\n"
         "하이어비자 모바일 서비스 이용을 진심으로 환영합니다. 😊\n"
         "이용 중 불편하신 사항은 언제든지 채팅으로 문의해 주세요.",
         "", "{이름}\n{개통유형}\n{SIM유형}", "채팅 종료됨", EMERALD_LT),

        ("활성화 확인\n완료 후",
         "미정의 (운영팀 확정 필요)",
         "(미정의) 개통 완료 후 학생이 실제 활성화를 확인했을 때",
         "(현재 없음 — 운영팀 정의 필요)",
         "", "(미정의)", "예: 'SIM 활성화가 정상 확인되었습니다' 형태", AMBER_LT),

        ("서류 보완 요청\n(항목별 템플릿)",
         "수동 발송 (즉시)\n오퍼레이터 SuppModal 선택 후",
         "오퍼레이터가 보완 요청 버튼 클릭 → 사유 선택",
         "[보완 요청] {서류명}\n{보완 사유} · {추가 안내(선택)}",
         "", "{서류명}\n{보완 사유}\n{추가 안내}", "항목별 보완 사유는 '📨 보완 요청 문구' 시트 참고", AMBER_LT),

        ("서류 제출 요청\n(미제출 항목)",
         "수동 발송 (즉시)\n오퍼레이터 클릭 후",
         "오퍼레이터가 미제출(empty) 항목에 '📋 제출 요청' 버튼 클릭",
         "[보완 요청] {서류명}\n서류 미제출",
         "", "{서류명}", "앱번호 경로에서 일부 서류 누락 시 발생", GRAY_LT),

        ("무응답 안내\n(대기 타이머)",
         "자동 발송 (60초)\n봇 발송",
         "학생 메시지/제출 후 60초간 오퍼레이터 무응답",
         "⏳ 곧 응대 예정입니다. 잠시만 기다려주세요.",
         "", "없음", "오퍼레이터 처리 시 타이머 리셋", ORANGE_LT),

        ("사증발급확인서\n요청",
         "수동 발송 (즉시)\n오퍼레이터 버튼 클릭",
         "체크리스트 '🛂 사증발급확인서 요청' 버튼 클릭",
         "[추가 서류 요청] 사증발급확인서\n자사에서 확보하지 못한 관계로 사증발급확인서 제출을 요청드립니다.",
         "", "없음", "자사 미확보 케이스에만 사용", PURPLE_LT),
    ]

    for r_idx, row in enumerate(templates, 4):
        step, timing, cond, current, official, var, note, bg = row
        xc(ws, r_idx, 1, step,    bg=bg, bold=True, align="center")
        xc(ws, r_idx, 2, timing,  bg=bg, align="center")
        xc(ws, r_idx, 3, cond,    bg=GRAY_LT)
        xc(ws, r_idx, 4, current, bg=WHITE, fg=GRAY, italic=True)
        xinput(ws, r_idx, 5, "(공식 문구 작성 — 미작성 시 임시 문구 사용)")
        xc(ws, r_idx, 6, var,  bg=GRAY_LT, fg=NAVY, bold=True)
        xc(ws, r_idx, 7, note, bg=WHITE, fg=GRAY, italic=True)
        ws.row_dimensions[r_idx].height = 80

    note_row(ws, len(templates) + 4 + 1,
             "💡 변수({이름} 등)는 시스템이 자동으로 치환합니다. "
             "공식 문구 작성 시 변수 자리({  })는 그대로 유지해 주세요.", ncols=7)
    input_legend(ws, len(templates) + 4 + 3, ncols=7)


# ══════════════════════════════════════════════════════════════════
# XLSX Sheet 5: 자동화 워크플로우
# ══════════════════════════════════════════════════════════════════
def build_automation(wb):
    ws = wb.create_sheet("⚙️ 자동화 워크플로우")
    sheet_banner(ws, "⚙️ 자동화 워크플로우 정의",
                 "오퍼레이터 개입 없이 자동 처리되는 흐름 확인 및 보완")

    headers = ["단계", "진행 주체", "발생 조건", "자동 처리 내용",
               "고객에게 보이는 것", "오퍼레이터에게 보이는 것",
               "확정 여부", "수정 요청 사항"]
    widths  = [18,     12,          32,           32,
               28,                   28,
               12,      28]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        xh(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    flows = [
        ("1. 신청번호 입력",
         "시스템", "학생이 신청번호 입력 후 '채팅 시작' 클릭",
         "① '신청번호 확인' 봇 메시지 발송\n② 1분 대기 타이머 시작\n③ 오퍼레이터 어드민에 Seg 확인 카드 노출",
         "'신청번호 확인' 메시지\n'잠시 후 오퍼레이터가 안내해 드릴게요'",
         "Seg 확인 카드 (S1/S2 × eSIM/uSIM 4가지 버튼)",
         "✅ 확정", "", BLUE_LT),

        ("2. Seg 확인\n(최초)",
         "오퍼레이터\n→ 자동",
         "오퍼레이터가 Seg 확인 카드에서 유형 선택",
         "① 해당 Seg 서류 체크리스트 활성화\n② 인트로 메시지 자동 발송 (300ms)\n③ 대기 타이머 중지",
         "인트로 메시지 (채팅에 표시)\n서류 제출 현황 패널 표시",
         "체크리스트 활성화\n항목별 승인/보완/제출요청 버튼",
         "⚠️ 확인 필요", "외등 Seg 추가 시 카드 항목 확장 필요", AMBER_LT),

        ("3. 서류 보완 요청",
         "오퍼레이터\n→ 자동",
         "오퍼레이터가 '📨 보완 요청' 버튼 → 사유 선택 → 전송",
         "① 보완 요청 템플릿 메시지 발송\n② 해당 항목 상태 → supp_requested\n③ Status → B (보완 요청 중)",
         "보완 요청 메시지 (주황 카드)\n서류 현황: '보완 요청' 표시\n퀵리플라이 버튼 3개 노출",
         "해당 항목 '학생 대기 중' 표시",
         "✅ 확정", "", AMBER_LT),

        ("4. 학생 재제출",
         "학생\n→ 자동",
         "학생이 보완 항목 재제출 클릭 → 파일/텍스트 업로드",
         "① 항목 상태 → resubmitted\n"
         "② 마지막 재제출 시 Status → C\n"
         "③ Status C 전환 시 '재제출 확인' 메시지 자동 발송 (600ms)",
         "(서류 업로드 UI)\n재제출됨 상태 배지",
         "항목 상태 '재검토 대기' 표시",
         "✅ 확정", "", PURPLE_LT),

        ("5. 전체 승인\n→ 개통 준비",
         "오퍼레이터\n→ 자동",
         "오퍼레이터가 마지막 항목 승인 → 전체 approved",
         "① Status → D (개통 준비 완료)\n② '전체 승인 완료' 메시지 자동 발송 (400ms)\n③ '개통 완료 처리' 버튼 활성화",
         "'모든 서류 확인 완료' 메시지\n개통 준비 완료 배지",
         "초록 '개통 완료 처리' 버튼 카드 표시",
         "✅ 확정", "", EMERALD_LT),

        ("6. 개통 완료 처리",
         "오퍼레이터\n→ 자동",
         "오퍼레이터가 '✅ 개통 완료 처리' 버튼 클릭",
         "① 개통 완료 메시지 발송\n② 채팅 종료 처리 (closed = true)",
         "개통 완료 메시지 (이름/유형 포함)\n채팅 종료 상태",
         "채팅 종료됨 배지",
         "✅ 확정", "", EMERALD_LT),

        ("7. 활성화 확인\n완료",
         "(미정의)",
         "(미정의) 개통 후 실제 SIM 활성화 확인 시점",
         "(미정의) 활성화 확인 메시지 발송 여부 및 방식 확정 필요",
         "(미정의)",
         "(미정의)",
         "❌ 미정의", "활성화 확인 시점 및 메시지 여부 운영팀 결정 필요", AMBER_LT),

        ("8. 무응답 안내\n(1분)",
         "시스템\n(봇)",
         "학생 메시지/제출 후 60초 무응답",
         "봇이 '곧 응대 예정입니다' 메시지 자동 발송",
         "안내 메시지 표시",
         "(없음)",
         "⚠️ 확인 필요", "60초 기준 적합한지 운영팀 확인 필요", ORANGE_LT),
    ]

    for r_idx, row in enumerate(flows, 4):
        step, actor, cond, auto, cust_view, op_view, confirmed, request, bg = row
        xc(ws, r_idx, 1, step,      bg=bg, bold=True, align="center")
        xc(ws, r_idx, 2, actor,     bg=bg, align="center")
        xc(ws, r_idx, 3, cond,      bg=GRAY_LT)
        xc(ws, r_idx, 4, auto,      bg=bg)
        xc(ws, r_idx, 5, cust_view, bg=BLUE_LT if "미정의" not in cust_view else AMBER_LT)
        xc(ws, r_idx, 6, op_view,   bg=GRAY_LT)
        # 확정 여부
        fg_conf = EMERALD if "확정" in confirmed else (RED if "미정의" in confirmed else AMBER)
        xc(ws, r_idx, 7, confirmed, bg=WHITE, fg=fg_conf, bold=True, align="center")
        # 수정 요청
        if request:
            xinput(ws, r_idx, 8, request)
        else:
            xc(ws, r_idx, 8, "", bg=WHITE)
        ws.row_dimensions[r_idx].height = 72

    note_row(ws, len(flows) + 5,
             "💡 '확정 여부'가 ⚠️ 또는 ❌인 항목은 운영팀 검토 후 '수정 요청 사항'에 의견을 기재해 주세요.", ncols=8)
    input_legend(ws, len(flows) + 7, ncols=8)


# ══════════════════════════════════════════════════════════════════
# XLSX Sheet 6: 보완 요청 표준 문구
# ══════════════════════════════════════════════════════════════════
def build_supp_scripts(wb):
    ws = wb.create_sheet("📨 보완 요청 표준 문구")
    sheet_banner(ws, "📨 보완 요청 표준 문구 작성",
                 "서류별 보완 요청 시 고객에게 발송할 표준 안내 문구")

    headers = ["서류명 (Seg)", "보완 사유 (선택지 텍스트)",
               "현재 발송 형식 (참고)", "추가 안내 표준 문구 (운영팀 작성)", "비고"]
    widths  = [24,              36,
               36,               50,                              16]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        xh(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    items = [
        # 여권
        ("여권 양면 평판 스캔본", "사진이 흐릿함 - 재스캔 필요",
         "[보완 요청] 여권 양면 평판 스캔본\n사진이 흐릿함 - 재스캔 필요",
         BLUE_LT),
        ("여권 양면 평판 스캔본", "한쪽 면만 보임 - 양면 모두 포함 필요",
         "[보완 요청] 여권 양면 평판 스캔본\n한쪽 면만 보임 - 양면 모두 포함 필요",
         BLUE_LT),
        ("여권 양면 평판 스캔본", "정보 일부 가림 - 정면 재스캔",
         "[보완 요청] 여권 양면 평판 스캔본\n정보 일부 가림 - 정면 재스캔",
         BLUE_LT),
        ("여권 양면 평판 스캔본", "만료일 임박/지남 - 재발급 후 제출",
         "[보완 요청] 여권 양면 평판 스캔본\n만료일 임박/지남 - 재발급 후 제출",
         BLUE_LT),
        # 기기정보
        ("기기정보 캡처본 (eSIM)", "IMEI가 보이지 않음 - *#06# 화면 재캡처",
         "[보완 요청] 기기정보 캡처본\nIMEI가 보이지 않음",
         PURPLE_LT),
        ("기기정보 캡처본 (eSIM)", "모델명 확인 불가 - 설정 > 일반 > 정보 화면 캡처",
         "[보완 요청] 기기정보 캡처본\n모델명 확인 불가",
         PURPLE_LT),
        ("기기정보 캡처본 (eSIM)", "eSIM 지원 여부 확인 불가 - eSIM 항목 포함 재캡처",
         "[보완 요청] 기기정보 캡처본\neSIM 지원 여부 확인 불가",
         PURPLE_LT),
        # 유심
        ("유심 사진 (uSIM)", "사진이 흐릿함 - 재촬영 필요",
         "[보완 요청] 유심 사진\n사진이 흐릿함 - 재촬영 필요",
         TEAL_LT),
        ("유심 사진 (uSIM)", "유심 번호가 보이지 않음 - 번호 부분 포함 재촬영",
         "[보완 요청] 유심 사진\n유심 번호가 보이지 않음",
         TEAL_LT),
        # 기존 통신사
        ("기존 통신사 확인 (S2)", "문자 캡처본이 불명확함 - 재캡처 필요",
         "[보완 요청] 기존 통신사 확인\n문자 캡처본이 불명확함",
         AMBER_LT),
        ("기존 통신사 확인 (S2)", "다른 번호의 캡처본으로 보임 - 본인 번호로 재확인",
         "[보완 요청] 기존 통신사 확인\n다른 번호 캡처본으로 보임",
         AMBER_LT),
        # 기존 전화번호
        ("기존 전화번호 (S2)", "번호 형식 오류 - 올바른 번호 재입력",
         "[보완 요청] 기존 전화번호\n번호 형식 오류",
         AMBER_LT),
        # 외등 (미정의)
        ("외국인등록증 (외등 Seg)", "(운영팀 보완 사유 추가 필요)",
         "(미정의)",
         EMERALD_LT),
    ]

    for r_idx, row in enumerate(items, 4):
        doc_name, reason, current, bg = row
        xc(ws, r_idx, 1, doc_name, bg=bg, bold=True)
        xc(ws, r_idx, 2, reason,   bg=bg)
        xc(ws, r_idx, 3, current,  bg=WHITE, fg=GRAY, italic=True)
        xinput(ws, r_idx, 4,
               "(이 보완 사유로 요청할 때 추가로 안내할 표준 문구 작성\n"
               "예: '여권 정보면과 비자 페이지가 모두 포함된 평판 스캔본으로 재제출 부탁드립니다.')")
        xc(ws, r_idx, 5, "", bg=WHITE)
        ws.row_dimensions[r_idx].height = 52

    note_row(ws, len(items) + 5,
             "💡 현재 발송 형식은 '[보완 요청] 서류명 + 사유' 형태입니다. "
             "추가 안내 문구는 이 뒤에 자동으로 이어 붙여집니다.", ncols=5)
    input_legend(ws, len(items) + 7, ncols=5)


# ══════════════════════════════════════════════════════════════════
# XLSX Sheet 7: 퀵리플라이 응답 스크립트
# ══════════════════════════════════════════════════════════════════
def build_quickreply_scripts(wb):
    ws = wb.create_sheet("❓ 퀵리플라이 응답 스크립트")
    sheet_banner(ws, "❓ 퀵리플라이 버튼 응답 스크립트",
                 "학생이 버튼을 눌렀을 때 오퍼레이터가 발송할 표준 응답")

    # 노출 조건 안내
    ws.row_dimensions[2].height = 8
    r = 3
    xh(ws, r, 1, "퀵리플라이 버튼이 표시되는 조건", bg=AMBER_LT, fg="92400E", colspan=6)
    ws.row_dimensions[r].height = 20
    r += 1
    c = ws.cell(row=r, column=1,
                value="보완 요청 중(Status B)이고, 아직 상담사 연결을 요청하지 않은 상태에서만 표시됩니다.\n"
                      "고객이 버튼을 누르면 해당 텍스트가 고객 메시지로 전송되고, 오퍼레이터는 이 스크립트를 기준으로 응답합니다.")
    c.fill = xfill(AMBER_LT)
    c.font = xfont(size=9)
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 38
    r += 2

    headers = ["버튼 텍스트", "고객 발송 메시지", "응답 방식",
               "오퍼레이터 응답 스크립트 (운영팀 작성)",
               "항목별 분기 여부", "비고"]
    widths  = [30,             30,               14,
               50,
               14,              20]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        xh(ws, r, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w
    r += 1

    buttons = [
        ("❓ 어떤 부분이 부족한 건지 모르겠어요",
         "어떤 부분이 부족한 건지 모르겠어요",
         "오퍼레이터 직접 응답\n(자유 채팅 미활성)",
         AMBER_LT),
        ("❓ 어떻게 제출해야 하는지 모르겠어요",
         "어떻게 제출해야 하는지 모르겠어요",
         "오퍼레이터 직접 응답\n(자유 채팅 미활성)",
         AMBER_LT),
        ("💬 상담사 연결이 필요해요",
         "상담사 연결이 필요해요",
         "자동: 시스템 알림 발송\n→ 오퍼레이터 자유 채팅 활성화",
         NAVY_LT),
    ]

    for row_idx, (btn_text, msg, resp_type, bg) in enumerate(buttons, r):
        xc(ws, row_idx, 1, btn_text,   bg=bg, bold=True)
        xc(ws, row_idx, 2, msg,        bg=bg)
        xc(ws, row_idx, 3, resp_type,  bg=bg, align="center")
        xinput(ws, row_idx, 4,
               f"('{btn_text}' 버튼 클릭 시 오퍼레이터 응답 스크립트 작성)")
        xinput(ws, row_idx, 5, "Y / N")
        xinput(ws, row_idx, 6, "항목별 분기 Y인 경우 아래 테이블에 항목별 작성")
        ws.row_dimensions[row_idx].height = 60
    r += len(buttons)

    # 항목별 분기 테이블
    r += 2
    xh(ws, r, 1, "항목별 분기 응답 (필요 시 작성)", bg=NAVY_LT, fg=NAVY, colspan=6)
    ws.row_dimensions[r].height = 20
    r += 1

    sub_headers = ["버튼", "해당 서류명", "항목별 응답 스크립트", "비고"]
    sub_widths  = [30,     22,            50,                    16]
    for col, (h, w) in enumerate(zip(sub_headers, sub_widths), 1):
        xh(ws, r, col, h, bg=GRAY_LT, fg="1F2937", size=9)
        ws.column_dimensions[get_column_letter(col)].width = w
    r += 1

    sub_rows = [
        ("❓ 어떤 부분이 부족한 건지 모르겠어요", "여권 양면 평판 스캔본", "", ""),
        ("", "기기정보 캡처본", "", ""),
        ("", "유심 사진", "", ""),
        ("", "기존 통신사 확인", "", ""),
        ("", "기존 전화번호", "", ""),
        ("❓ 어떻게 제출해야 하는지 모르겠어요", "여권 양면 평판 스캔본", "", ""),
        ("", "기기정보 캡처본", "", ""),
        ("", "유심 사진", "", ""),
        ("", "기존 통신사 확인", "", ""),
        ("", "기존 전화번호", "", ""),
    ]
    for sub_row in sub_rows:
        xc(ws, r, 1, sub_row[0], bg=GRAY_LT if sub_row[0] else WHITE)
        xc(ws, r, 2, sub_row[1], bg=GRAY_LT)
        xinput(ws, r, 3, "(항목별 응답 작성)")
        xc(ws, r, 4, "", bg=WHITE)
        ws.row_dimensions[r].height = 36
        r += 1

    input_legend(ws, r + 1, ncols=6)


# ══════════════════════════════════════════════════════════════════
# DOCX 헬퍼
# ══════════════════════════════════════════════════════════════════
def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_table_border(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), 'CCCCCC')
        borders.append(b)
    tblPr.append(borders)

def dp(doc, text, style='Normal', bold=False, size=None, color=None,
       space_before=0, space_after=4):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = 'Malgun Gothic'
    return p

def dh1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    run.font.name = 'Malgun Gothic'
    # Bottom border
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:color'), '1F4E79')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def dh2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    run.font.name = 'Malgun Gothic'
    return p

def dh3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string("374151")
    run.font.name = 'Malgun Gothic'
    return p

def dbullet(doc, items, level=0):
    indent = "   " * level + "• "
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0.5 + level * 0.5)
        run = p.add_run(indent + item)
        run.font.size = Pt(9.5)
        run.font.name = 'Malgun Gothic'

def dtable(doc, headers, rows, col_widths=None, header_bg=NAVY):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    add_table_border(table)

    # Header row
    hrow = table.rows[0]
    hrow.height = Cm(0.8)
    for i, (h, cell) in enumerate(zip(headers, hrow.cells)):
        set_cell_bg(cell, header_bg)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.name = 'Malgun Gothic'
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Data rows
    for r_idx, row in enumerate(rows):
        drow = table.rows[r_idx + 1]
        for c_idx, (val, cell) in enumerate(zip(row, drow.cells)):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            run.font.name = 'Malgun Gothic'

    if col_widths:
        for row in table.rows:
            for i, (cell, w) in enumerate(zip(row.cells, col_widths)):
                cell.width = Cm(w)

    return table

def dbox(doc, title, body_lines, bg=NAVY_LT, title_color=NAVY):
    """Simple bordered notice box using a 1×1 table"""
    table = doc.add_table(rows=1, cols=1)
    add_table_border(table)
    cell = table.rows[0].cells[0]
    set_cell_bg(cell, bg)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title + "\n")
    run.bold = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(title_color)
    run.font.name = 'Malgun Gothic'
    for line in body_lines:
        run2 = p.add_run("   " + line + "\n")
        run2.font.size = Pt(9)
        run2.font.name = 'Malgun Gothic'
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ══════════════════════════════════════════════════════════════════
# DOCX 생성
# ══════════════════════════════════════════════════════════════════
def build_docx():
    doc = Document()

    # 페이지 여백 설정
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(2.5)

    # ── 표지 헤더 ───────────────────────────────────────────────
    # 수신/발신 블록
    meta_table = doc.add_table(rows=1, cols=2)
    meta_table.style = 'Table Grid'
    add_table_border(meta_table)
    left_cell  = meta_table.rows[0].cells[0]
    right_cell = meta_table.rows[0].cells[1]
    set_cell_bg(left_cell, NAVY)
    set_cell_bg(right_cell, NAVY_LT)

    lp = left_cell.paragraphs[0]
    lr = lp.add_run("HireVisa Mobile\nCS 채팅 시스템 운영 기준 정의 요청서")
    lr.bold = True
    lr.font.size = Pt(14)
    lr.font.color.rgb = RGBColor(255, 255, 255)
    lr.font.name = 'Malgun Gothic'
    lp.paragraph_format.space_before = Pt(4)
    lp.paragraph_format.space_after  = Pt(4)

    rp = right_cell.paragraphs[0]
    meta_lines = [
        ("수신", "운영팀"),
        ("발신", "개발/기획팀"),
        ("작성일", "2026-05-26"),
        ("버전", "v1.0"),
    ]
    for label, val in meta_lines:
        run = rp.add_run(f"{label}: ")
        run.bold = True
        run.font.size = Pt(9)
        run.font.name = 'Malgun Gothic'
        run2 = rp.add_run(f"{val}\n")
        run2.font.size = Pt(9)
        run2.font.name = 'Malgun Gothic'

    left_cell.width  = Cm(10)
    right_cell.width = Cm(6.5)
    doc.add_paragraph()

    # ── 1. 요청 배경 ───────────────────────────────────────────
    dh1(doc, "1. 요청 배경 및 목적")

    dp(doc,
       "HireVisa Mobile CS 채팅 시스템(이하 '채팅 시스템')이 Beta 단계로 개발 완료되었습니다. "
       "본 시스템은 유학생의 한국 모바일 개통 신청부터 서류 확인, 개통 완료까지의 전 과정을 "
       "채팅 인터페이스로 지원합니다.",
       size=10)

    dp(doc,
       "정식 운영을 위해 아래 3가지 사항에 대해 운영팀의 공식 기준 정의가 필요합니다. "
       "첨부된 작성 양식(ops-request.xlsx)에 운영팀의 의견을 기재하여 회신해 주시기 바랍니다.",
       size=10)

    dbullet(doc, [
        "Seg(개통 유형)별 응대 시나리오 정의",
        "단계별 채팅 템플릿 메시지 작성",
        "자동화 워크플로우 기준 확인 및 보완",
    ])

    # ── 2. 시스템 개요 ─────────────────────────────────────────
    dh1(doc, "2. 채팅 시스템 개요 (비기술적 설명)")

    dh2(doc, "2.1 학생 진입 경로")
    dp(doc, "학생이 채팅에 진입하는 방법은 두 가지입니다:", size=10)
    dbullet(doc, [
        "경로 A — 신청번호 입력: mobile.hirevisa.com/en에서 신청 완료 후 발급받은 번호로 채팅 시작. "
          "시스템은 모든 서류가 이미 제출되었다고 가정하여 오퍼레이터가 즉시 검토 가능 상태로 설정됩니다.",
        "경로 B — 채팅 내 간편신청: 신청번호가 없는 학생이 채팅에서 개통 유형·SIM 유형을 선택하여 신청번호를 발급받고, "
          "이후 채팅으로 서류를 직접 제출합니다.",
    ])

    dh2(doc, "2.2 Seg (개통 유형) 구분")
    dp(doc, "Seg는 신분증 유형(여권/외등)과 개통 방식(신규/번이)의 조합으로 총 4가지입니다:", size=10)

    dtable(doc,
           ["Seg명", "신분증", "개통 방식", "주요 대상"],
           [
               ["여권 신규", "여권 (Passport)", "신규 번호 발급", "한국 첫 입국 유학생"],
               ["여권 번이", "여권 (Passport)", "기존 번호 이동 (MNP)", "여권 소지 + 기존 한국 번호 보유"],
               ["외등 신규", "외국인등록증 (ARC)", "신규 번호 발급", "장기 거주 유학생 (90일 이상)"],
               ["외등 번이", "외국인등록증 (ARC)", "기존 번호 이동 (MNP)", "장기 거주 + 기존 번호 보유"],
           ],
           col_widths=[3, 4, 4, 5.5],
           header_bg=NAVY)
    doc.add_paragraph()

    dbox(doc,
         "⚠️ 외등(외국인등록증) Seg 관련 안내",
         ["현재 시스템(Beta)은 여권 Seg만 구현되어 있습니다.",
          "외등 Seg는 운영팀이 필요 서류 기준을 확정하면 개발팀이 시스템에 반영할 예정입니다.",
          "외등 관련 항목(📄 필요 서류 정의 시트 노란 셀)을 반드시 작성해 주세요."],
         bg=AMBER_LT, title_color=ORANGE)

    dh2(doc, "2.3 서류 처리 흐름")
    dp(doc, "채팅 내 서류 처리는 아래 순서로 진행됩니다:", size=10)
    dbullet(doc, [
        "오퍼레이터가 Seg 확인 → 서류 목록 활성화",
        "오퍼레이터가 각 서류를 검토 후 ✅ 승인 또는 📨 보완 요청",
        "보완 요청 시 학생에게 메시지 자동 발송 → 학생 재제출",
        "모든 서류 승인 완료(Status D) → 개통 완료 처리 버튼 활성화",
        "오퍼레이터 클릭 시 개통 완료 메시지 자동 발송 → 채팅 종료",
    ])

    dh2(doc, "2.4 자동 발송 메시지")
    dp(doc,
       "오퍼레이터가 별도로 메시지를 작성하지 않아도, 특정 시점에 미리 정의된 문구가 자동으로 발송됩니다. "
       "현재는 임시 문구가 설정되어 있으며, 이번 요청을 통해 공식 문구로 교체됩니다.",
       size=10)

    # ── 3. 업무 요청 ───────────────────────────────────────────
    dh1(doc, "3. 업무 요청 사항")

    # 3.1
    dh2(doc, "요청 1. Seg별 응대 시나리오 정의")
    dp(doc, "▸ 첨부 파일 작성 위치: [🗂 Seg별 응대 시나리오] 시트", bold=True, size=9,
       color=NAVY)
    dp(doc,
       "4개 Seg 각각에 대해 오퍼레이터가 따라야 할 응대 흐름을 정의해 주세요. "
       "시스템이 자동으로 처리하는 단계는 이미 구현되어 있으나, "
       "오퍼레이터 판단이 필요한 구간(서류 검토 기준, 이슈 처리 방법 등)은 운영 매뉴얼 기준으로 작성이 필요합니다.",
       size=10)
    dh3(doc, "작성 항목:")
    dbullet(doc, [
        "응대 순서 (단계별 정리): 오퍼레이터가 실제로 수행하는 행동 순서",
        "특이사항 / 자주 발생 이슈: 현장에서 자주 발생하는 예외 케이스",
        "예상 처리 시간: Seg별 평균 처리 소요 시간",
        "오퍼레이터 주의사항: 실수하기 쉬운 체크포인트",
    ])

    # 3.2
    dh2(doc, "요청 2. 채팅 템플릿 메시지 작성")
    dp(doc, "▸ 첨부 파일 작성 위치: [💬 채팅 템플릿 메시지] 시트 + [📨 보완 요청 표준 문구] 시트",
       bold=True, size=9, color=NAVY)
    dp(doc,
       "아래 5개 단계에서 고객에게 발송될 공식 문구를 작성해 주세요. "
       "각 시트에 현재 임시 문구(회색)가 참고용으로 표시되어 있으며, "
       "노란색 셀에 공식 문구를 작성하시면 됩니다.",
       size=10)

    dtable(doc,
           ["단계", "발송 시점", "현재 상태", "작성 필요"],
           [
               ["채팅 시작 인트로", "Seg 확인 직후 자동 발송", "임시 문구 있음", "공식 문구 교체"],
               ["서류 검토 중 (재제출 확인)", "학생 재제출 후 자동 발송", "임시 문구 있음", "공식 문구 교체"],
               ["개통 전 (전체 승인)", "마지막 서류 승인 후 자동 발송", "임시 문구 있음", "공식 문구 교체"],
               ["개통 완료 후", "개통 완료 처리 클릭 시 자동 발송", "임시 문구 있음", "공식 문구 교체"],
               ["활성화 확인 완료 후", "미정의", "없음", "발송 여부 및 문구 신규 작성"],
           ],
           col_widths=[4, 5, 4, 4.5],
           header_bg=NAVY)
    doc.add_paragraph()

    dp(doc, "서류 보완 요청 시 각 항목별 표준 안내 문구도 함께 작성이 필요합니다 "
            "([📨 보완 요청 표준 문구] 시트).", size=10)

    # 3.3
    dh2(doc, "요청 3. 자동화 워크플로우 확인 및 보완")
    dp(doc, "▸ 첨부 파일 작성 위치: [⚙️ 자동화 워크플로우] 시트 + [❓ 퀵리플라이 응답 스크립트] 시트",
       bold=True, size=9, color=NAVY)
    dp(doc,
       "시스템에 현재 구현된 자동화 흐름을 검토하고, 운영 기준과 맞지 않는 부분을 알려주세요. "
       "특히 아래 항목은 운영팀의 확인이 반드시 필요합니다:",
       size=10)
    dbullet(doc, [
        "무응답 대기 시간 (현재 60초): 실제 운영 환경에 맞는 기준 확인",
        "활성화 확인 완료 단계: 자동화 메시지 발송 여부 및 시점 결정",
        "퀵리플라이 버튼 응답 스크립트: "
          "학생이 '어떤 부분이 부족한지 모르겠어요' 등을 클릭했을 때 오퍼레이터가 참고할 표준 응답 작성",
        "외등 Seg 추가 시 Seg 확인 카드 및 워크플로우 변경사항 확인",
    ])

    # ── 4. 첨부 파일 활용 안내 ─────────────────────────────────
    dh1(doc, "4. 첨부 파일 활용 안내")

    dtable(doc,
           ["파일명", "용도", "비고"],
           [
               ["ops-request.xlsx", "운영팀 작성 양식 (본 요청서 작성 대상)", "노란색 셀에 직접 작성"],
               ["workflow-spec.md", "시스템 전체 워크플로우 기술 명세 (참고용)", "개발팀 작성, 참고용"],
               ["workflow-branches.xlsx", "기술 명세 스프레드시트 (참고용)", "개발팀 작성, 참고용"],
           ],
           col_widths=[5, 7, 4.5],
           header_bg=NAVY)
    doc.add_paragraph()

    dbox(doc,
         "📝 작성 방법",
         ["ops-request.xlsx를 열어 [📌 안내 및 체크리스트] 시트부터 순서대로 읽고 시작해 주세요.",
          "노란색(🟡) 셀만 작성하면 됩니다. 흰색 셀은 참고용입니다.",
          "잘 모르는 항목은 비워두고 '비고'에 질문을 기재해 주시면 협의하겠습니다.",
          "작성 완료 후 ops-request.xlsx 파일을 이메일/슬랙으로 회신해 주세요."],
         bg=EMERALD_LT, title_color=EMERALD)

    # ── 5. 제출 기한 및 협의 ───────────────────────────────────
    dh1(doc, "5. 제출 기한 및 협의 사항")

    dtable(doc,
           ["항목", "내용"],
           [
               ["작성 기한", "(협의 후 기재)"],
               ["제출 방법", "작성 완료된 ops-request.xlsx 파일 회신"],
               ["담당자 문의", "(개발/기획 담당자 연락처)"],
               ["협의 방법", "불명확한 항목은 회의 통해 협의 가능"],
           ],
           col_widths=[4.5, 12],
           header_bg=NAVY)
    doc.add_paragraph()

    dbox(doc,
         "⚠️ 우선 작성 항목 (필수)",
         ["외등(외국인등록증) Seg 필요 서류 목록 및 확인 기준 → [📄 필요 서류 정의] 시트",
          "개통 완료 후 공식 메시지 → [💬 채팅 템플릿 메시지] 시트",
          "활성화 확인 완료 후 메시지 여부 결정 → [💬 채팅 템플릿 메시지] + [⚙️ 자동화 워크플로우] 시트"],
         bg=RED_LT, title_color=RED)

    doc.add_paragraph()
    dp(doc, "— 이하 여백 —", size=9, color=GRAY)

    return doc


# ══════════════════════════════════════════════════════════════════
# Main
# ══════════════════════════════════════════════════════════════════
def main():
    # ── DOCX ──
    doc = build_docx()
    docx_path = "/home/user/cs_mobile_chat/docs/ops-request.docx"
    doc.save(docx_path)
    print(f"✅ DOCX saved: {docx_path}")

    # ── XLSX ──
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    build_guide(wb)
    build_seg_scenario(wb)
    build_docs_matrix(wb)
    build_templates(wb)
    build_automation(wb)
    build_supp_scripts(wb)
    build_quickreply_scripts(wb)

    xlsx_path = "/home/user/cs_mobile_chat/docs/ops-request.xlsx"
    wb.save(xlsx_path)
    print(f"✅ XLSX saved: {xlsx_path}")
    print(f"   Sheets: {[ws.title for ws in wb.worksheets]}")


if __name__ == "__main__":
    main()
