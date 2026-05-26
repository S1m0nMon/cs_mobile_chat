#!/usr/bin/env python3
"""
Generate workflow-branches.xlsx — HireVisa Mobile CS Chat
Multi-sheet workbook covering all branching logic derived from the codebase.
Run: python3 docs/generate_xlsx.py
"""

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

# ── Color palette ──────────────────────────────────────────────────────────────
NAVY        = "1F4E79"
NAVY_LIGHT  = "D6E4F0"
NAVY_MID    = "A8C7E2"
AMBER       = "F59E0B"
AMBER_LIGHT = "FEF3C7"
EMERALD     = "059669"
EMERALD_LT  = "D1FAE5"
PURPLE      = "7C3AED"
PURPLE_LT   = "EDE9FE"
BLUE        = "2563EB"
BLUE_LT     = "DBEAFE"
RED         = "DC2626"
RED_LT      = "FEE2E2"
GRAY        = "6B7280"
GRAY_LT     = "F3F4F6"
GRAY_MID    = "E5E7EB"
WHITE       = "FFFFFF"
ORANGE_LT   = "FFF7ED"
ORANGE      = "EA580C"
TEAL        = "0D9488"
TEAL_LT     = "CCFBF1"

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def font(bold=False, color="000000", size=10, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic,
                name="Malgun Gothic")

def border_all(thin=True):
    s = Side(style="thin" if thin else "hair", color="D1D5DB")
    return Border(left=s, right=s, top=s, bottom=s)

def center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def make_header(ws, row, col, text, bg=NAVY, fg=WHITE, bold=True, size=10):
    c = ws.cell(row=row, column=col, value=text)
    c.fill = fill(bg)
    c.font = font(bold=bold, color=fg, size=size)
    c.alignment = center()
    c.border = border_all()
    return c

def make_cell(ws, row, col, text, bg=WHITE, fg="1F2937", bold=False,
              align="left", size=10):
    c = ws.cell(row=row, column=col, value=text)
    c.fill = fill(bg)
    c.font = font(bold=bold, color=fg, size=size)
    c.alignment = left() if align == "left" else center()
    c.border = border_all()
    return c

def sheet_title(ws, title, subtitle=""):
    ws.row_dimensions[1].height = 32
    c = ws.cell(row=1, column=1, value=title + (f"  ·  {subtitle}" if subtitle else ""))
    c.fill = fill(NAVY)
    c.font = font(bold=True, color=WHITE, size=13)
    c.alignment = center()
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=20)
    ws.row_dimensions[2].height = 6  # spacer


# ══════════════════════════════════════════════════════════════════════════════
# 1. 진입 분기 (Entry Path)
# ══════════════════════════════════════════════════════════════════════════════
def build_entry_path(wb):
    ws = wb.create_sheet("1. 진입 분기")
    sheet_title(ws, "1. 진입 분기 (Entry Path)", "학생이 채팅에 진입하는 두 가지 경로")

    headers = ["#", "조건", "경로명", "학생 첫 화면", "항목 초기 상태",
               "오퍼레이터 첫 동작", "Seg 결정 주체", "인트로 딜레이"]
    widths =  [4,   28,     14,       32,             20,
               28,                    18,              16]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        make_header(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    rows = [
        ["A", "신청번호 보유\n(mobile.hirevisa.com/en 신청 완료)",
         "앱번호 경로",
         "신청번호 입력창\n(HV2026-XXXXX 형식)",
         "submitted\n(서류 이미 제출됨 가정)",
         "Seg 확인 카드 표시\n→ S1/S2 × eSIM/uSIM 4가지 중 선택",
         "오퍼레이터",
         "300ms (Seg 확인 후)"],
        ["B", "신청번호 없음\n(미신청 또는 번호 분실)",
         "워크플로우 경로",
         "Q1: 개통 유형 선택\n→ Q2: SIM 선택\n→ Q3: 확인 & 제출",
         "empty\n(채팅으로 직접 제출)",
         "자동 인트로 발송\n(Seg 확인 카드 없음)",
         "학생 (Q1·Q2 선택)",
         "1800ms (Q3 완료 후)"],
    ]

    row_bgs = [NAVY_LIGHT, AMBER_LIGHT]
    for r_idx, row in enumerate(rows, 4):
        bg = row_bgs[r_idx - 4]
        make_cell(ws, r_idx, 1, row[0], bg=bg, bold=True, align="center")
        for c_idx, val in enumerate(row[1:], 2):
            make_cell(ws, r_idx, c_idx, val, bg=bg)
        ws.row_dimensions[r_idx].height = 60

    # Notes box
    ws.row_dimensions[7].height = 10
    note = ws.cell(row=8, column=1,
                   value="💡 핵심 차이: 앱번호 경로 항목은 'submitted'로 시작(오퍼레이터가 승인/보완만 처리), "
                         "워크플로우 경로는 'empty'로 시작(학생이 채팅으로 직접 제출).")
    note.fill = fill(GRAY_LT)
    note.font = font(italic=True, color=GRAY, size=9)
    note.alignment = left()
    ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=len(headers))
    ws.row_dimensions[8].height = 28


# ══════════════════════════════════════════════════════════════════════════════
# 2. 워크플로우 분기 (In-Chat Workflow)
# ══════════════════════════════════════════════════════════════════════════════
def build_workflow(wb):
    ws = wb.create_sheet("2. 워크플로우 분기")
    sheet_title(ws, "2. 워크플로우 분기 (In-Chat Application Q1→Q2→Q3)",
                "신청번호 없는 학생 전용")

    headers = ["단계", "질문", "선택지 텍스트", "아이콘", "서브텍스트",
               "내부 값", "다음 단계", "결정 항목"]
    widths =  [10,     28,     28,           8,      32,
               12,       12,       16]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        make_header(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    rows = [
        # Q1
        ["Q1", "어떤 개통을 원하시나요?",
         "신규 번호 개통", "🆕", "한국에서 새 번호를 발급받습니다.",
         "seg = S1", "Q2", "Seg"],
        ["Q1", "",
         "기존 번호 이동 (MNP)", "🔄", "사용 중인 번호를 그대로 가져옵니다.",
         "seg = S2", "Q2", "Seg"],
        # Q2
        ["Q2", "어떤 SIM을 사용하시나요?",
         "eSIM (디지털 SIM)", "📲", "기기에 내장된 SIM.\n기기정보 캡처본이 필요해요.",
         "simType = esim", "Q3", "SimType"],
        ["Q2", "",
         "uSIM (실물 SIM)", "💳", "물리 SIM 카드.\n유심 사진이 필요해요.",
         "simType = usim", "Q3", "SimType"],
        # Q3
        ["Q3", "신청 내용을 확인해주세요\n(선택 내역 요약 카드 표시)",
         "신청 완료 · 채팅 시작 →", "✅", "선택 내용 확인 후 채팅 시작",
         "onComplete(seg, simType)", "채팅 시작", "최종 확인"],
        ["Q3", "",
         "← 이전 단계", "↩️", "SIM 선택으로 돌아가기",
         "step = 'sim'", "Q2", "—"],
        # Cancel
        ["진입 화면\n(공통)", "이미 신청번호가 있어요\n(하단 링크)",
         "이미 신청번호가 있어요", "🔙", "워크플로우 취소 → 신청번호 입력 화면으로",
         "onCancel()", "앱번호 입력", "—"],
    ]

    row_bgs = {
        "Q1": BLUE_LT,
        "Q2": PURPLE_LT,
        "Q3": EMERALD_LT,
        "진입 화면\n(공통)": GRAY_LT,
    }

    for r_idx, row in enumerate(rows, 4):
        bg = row_bgs.get(row[0], WHITE)
        bold_step = row[0] in ("Q1", "Q2", "Q3")
        make_cell(ws, r_idx, 1, row[0], bg=bg, bold=bold_step, align="center")
        for c_idx, val in enumerate(row[1:], 2):
            make_cell(ws, r_idx, c_idx, val, bg=bg)
        ws.row_dimensions[r_idx].height = 50

    ws.row_dimensions[11].height = 10
    note = ws.cell(row=12, column=1,
                   value="💡 Q3 완료 시: 신청번호 자동 발급 (HV2026-XXXXX) → 항목 state='empty' 생성 → 1800ms 후 오퍼레이터 인트로 자동 발송")
    note.fill = fill(GRAY_LT)
    note.font = font(italic=True, color=GRAY, size=9)
    note.alignment = left()
    ws.merge_cells(start_row=12, start_column=1, end_row=12, end_column=len(headers))
    ws.row_dimensions[12].height = 22


# ══════════════════════════════════════════════════════════════════════════════
# 3. Seg별 항목 체크리스트
# ══════════════════════════════════════════════════════════════════════════════
def build_item_matrix(wb):
    ws = wb.create_sheet("3. 항목 체크리스트")
    sheet_title(ws, "3. Seg × SimType별 항목 체크리스트", "config/segments.json 기반")

    headers = ["Seg", "SimType", "순서", "항목 코드", "항목명", "아이콘",
               "타입", "힌트", "초기 상태\n(앱번호 경로)", "초기 상태\n(워크플로우)",
               "비고"]
    widths =  [8,     10,       6,    18,           20,    8,
               8,     36,       16,                  16,
               24]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        make_header(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    items = [
        # S1 eSIM
        ["S1", "eSIM", 1, "passport_scan",   "여권 양면 평판 스캔본",   "📘", "file",
         "여권 정보면+비자면이 모두 보이는 평판 스캔 1장",
         "submitted", "empty", "전 Seg 공통"],
        ["S1", "eSIM", 2, "device_info",      "기기정보 캡처본",          "📱", "file",
         "기기 설정 또는 *#06# 화면 캡처 (IMEI 포함)",
         "submitted", "empty", "eSIM 전용"],
        # S1 uSIM
        ["S1", "uSIM", 1, "passport_scan",   "여권 양면 평판 스캔본",   "📘", "file",
         "여권 정보면+비자면이 모두 보이는 평판 스캔 1장",
         "submitted", "empty", "전 Seg 공통"],
        ["S1", "uSIM", 2, "sim_photo",       "유심 사진",               "💳", "file",
         "유심 앞면이 명확히 보이는 사진",
         "submitted", "empty", "uSIM 전용"],
        # S2 eSIM
        ["S2", "eSIM", 1, "passport_scan",   "여권 양면 평판 스캔본",   "📘", "file",
         "여권 정보면+비자면이 모두 보이는 평판 스캔 1장",
         "submitted", "empty", "전 Seg 공통"],
        ["S2", "eSIM", 2, "current_carrier", "기존 통신사 확인",        "🏢", "file",
         "114에 문자 발송 후 수신된 통신사 안내 문자 캡처본",
         "submitted", "empty", "S2 전용"],
        ["S2", "eSIM", 3, "current_number",  "기존 전화번호",            "📞", "text",
         "번호이동할 현재 사용 중인 한국 휴대폰 번호\n(형식: 010-XXXX-XXXX)",
         "submitted", "empty", "S2 전용 · 텍스트 입력"],
        ["S2", "eSIM", 4, "device_info",     "기기정보 캡처본",          "📱", "file",
         "기기 설정 또는 *#06# 화면 캡처 (IMEI 포함)",
         "submitted", "empty", "eSIM 전용"],
        # S2 uSIM
        ["S2", "uSIM", 1, "passport_scan",   "여권 양면 평판 스캔본",   "📘", "file",
         "여권 정보면+비자면이 모두 보이는 평판 스캔 1장",
         "submitted", "empty", "전 Seg 공통"],
        ["S2", "uSIM", 2, "current_carrier", "기존 통신사 확인",        "🏢", "file",
         "114에 문자 발송 후 수신된 통신사 안내 문자 캡처본",
         "submitted", "empty", "S2 전용"],
        ["S2", "uSIM", 3, "current_number",  "기존 전화번호",            "📞", "text",
         "번호이동할 현재 사용 중인 한국 휴대폰 번호\n(형식: 010-XXXX-XXXX)",
         "submitted", "empty", "S2 전용 · 텍스트 입력"],
        ["S2", "uSIM", 4, "sim_photo",       "유심 사진",               "💳", "file",
         "유심 앞면이 명확히 보이는 사진",
         "submitted", "empty", "uSIM 전용"],
        # Optional
        ["(선택)", "전체", "—", "visa_confirm",   "사증발급확인서",          "🛂", "file",
         "법무부 발급 사증발급확인서 (자사 미확보 시 요청)",
         "N/A", "N/A", "오퍼레이터 수동 추가\n(상태: supp_requested)"],
    ]

    seg_colors = {
        "S1": BLUE_LT,
        "S2": AMBER_LIGHT,
        "(선택)": PURPLE_LT,
    }
    sim_colors = {
        "eSIM": "EDE9FE",
        "uSIM": TEAL_LT,
        "전체": GRAY_LT,
    }

    for r_idx, row in enumerate(items, 4):
        seg = row[0]
        sim = row[1]
        bg = seg_colors.get(seg, WHITE)
        make_cell(ws, r_idx, 1, seg,  bg=bg, bold=True, align="center")
        make_cell(ws, r_idx, 2, sim,  bg=sim_colors.get(sim, WHITE), align="center")
        make_cell(ws, r_idx, 3, row[2], bg=bg, align="center")
        make_cell(ws, r_idx, 4, row[3], bg=GRAY_LT)  # code
        make_cell(ws, r_idx, 5, row[4], bg=bg, bold=True)
        make_cell(ws, r_idx, 6, row[5], bg=bg, align="center", size=14)
        make_cell(ws, r_idx, 7, row[6], bg=bg, align="center")
        make_cell(ws, r_idx, 8, row[7], bg=bg)
        # initial state: app-number path
        state_a = row[8]
        bg_a = BLUE_LT if state_a == "submitted" else GRAY_LT if state_a == "N/A" else WHITE
        make_cell(ws, r_idx, 9, state_a, bg=bg_a, align="center",
                  fg=BLUE if state_a == "submitted" else GRAY)
        # initial state: workflow path
        state_w = row[9]
        bg_w = GRAY_LT if state_w in ("empty", "N/A") else WHITE
        make_cell(ws, r_idx, 10, state_w, bg=bg_w, align="center",
                  fg=GRAY)
        make_cell(ws, r_idx, 11, row[10], bg=WHITE)
        ws.row_dimensions[r_idx].height = 38


# ══════════════════════════════════════════════════════════════════════════════
# 4. 항목 상태 머신 (Item State Machine)
# ══════════════════════════════════════════════════════════════════════════════
def build_item_states(wb):
    ws = wb.create_sheet("4. 항목 상태 머신")
    sheet_title(ws, "4. 항목 상태 머신 (Item State Machine)", "empty → submitted → … → approved")

    headers = ["현재 상태", "레이블(오퍼레이터)", "레이블(학생)",
               "트리거 액션 (주체)", "다음 상태",
               "오퍼레이터 화면 색상", "학생 화면 색상",
               "오퍼레이터 가용 버튼"]
    widths = [16, 16, 14, 30, 16, 18, 18, 36]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        make_header(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    rows = [
        ["empty",         "미제출",       "미제출",
         "① 학생: 제출 버튼 클릭 → submitted\n② 오퍼레이터: 📋 제출 요청 → supp_requested",
         "submitted / supp_requested",
         "회색 (bg-gray-100)",  "회색 (bg-gray-100)",
         "📋 제출 요청"],
        ["submitted",     "검토 대기",    "제출됨",
         "① 오퍼레이터: ✅ 승인 → approved\n② 오퍼레이터: 📨 보완 요청 → supp_requested",
         "approved / supp_requested",
         "파랑 (bg-blue-100)", "파랑 (bg-blue-100)",
         "✅ 승인 / 📨 보완 요청"],
        ["supp_requested","학생 대기",    "보완 요청",
         "학생: 재제출 버튼 클릭 → resubmitted",
         "resubmitted",
         "주황 (bg-amber-100)", "주황 (bg-amber-100)\n+ 보완 사유 표시",
         "(없음 — 학생 대기 중)"],
        ["resubmitted",   "재검토 대기",  "재제출됨",
         "① 오퍼레이터: ✅ 승인 → approved\n② 오퍼레이터: 📨 추가 보완 → supp_requested",
         "approved / supp_requested",
         "보라 (bg-purple-100)", "보라 (bg-purple-100)",
         "✅ 승인 / 📨 추가 보완"],
        ["approved",      "✓ 승인",       "✓ 승인",
         "(없음 — 잠금 상태)",
         "—",
         "초록 (bg-emerald-100)", "초록 (bg-emerald-100)",
         "(잠금 — 변경 불가)"],
    ]

    state_bgs = {
        "empty":          GRAY_LT,
        "submitted":      BLUE_LT,
        "supp_requested": AMBER_LIGHT,
        "resubmitted":    PURPLE_LT,
        "approved":       EMERALD_LT,
    }

    for r_idx, row in enumerate(rows, 4):
        bg = state_bgs[row[0]]
        make_cell(ws, r_idx, 1, row[0], bg=bg, bold=True, align="center")
        for c_idx, val in enumerate(row[1:], 2):
            make_cell(ws, r_idx, c_idx, val, bg=bg)
        ws.row_dimensions[r_idx].height = 52


# ══════════════════════════════════════════════════════════════════════════════
# 5. Status 계산 (A/B/C/D)
# ══════════════════════════════════════════════════════════════════════════════
def build_status(wb):
    ws = wb.create_sheet("5. Status 계산 (A·B·C·D)")
    sheet_title(ws, "5. Case Status 자동 계산 (A/B/C/D)",
                "computeCaseStatus() — 항목 상태 기반 실시간 파생")

    headers = ["Status", "레이블", "계산 조건", "우선순위\n(높을수록 먼저)",
               "오퍼레이터 화면", "학생 배지", "학생 퀵리플라이", "다음 오퍼레이터 액션"]
    widths = [10, 14, 36, 10, 28, 16, 16, 28]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        make_header(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    rows = [
        ["null", "초기 (항목 없음)",
         "items.length === 0",
         "—",
         "체크리스트 비활성\n(Seg 확인 대기 또는 신청번호 입력 대기)",
         "검토대기 (회색)",
         "없음",
         "Seg 확인 카드 노출 (앱번호 경로)"],
        ["A", "검토 중",
         "기본값\n(D·B·C 어느 조건도 미해당)",
         "4 (최저)",
         "항목별 승인/보완요청/제출요청 버튼 활성",
         "A · 검토 중 (파랑)",
         "없음",
         "항목 하나씩 승인 또는 보완 요청"],
        ["B", "보완 요청 중",
         "items 중 1개 이상 state === 'supp_requested'",
         "2",
         "해당 항목: '학생 대기 중' 표시\n다른 항목은 계속 처리 가능",
         "B · 보완 요청 중 (주황)\n+ 보완 항목 목록",
         "✅ 3개 버튼 노출\n(부족한 부분/제출 방법/상담사)",
         "학생 재제출 대기\n(재제출 후 승인 또는 추가 보완)"],
        ["C", "재제출됨",
         "items 중 1개 이상 state === 'resubmitted'\n(B 조건 없음)",
         "3",
         "재검토 대기 항목 승인/추가 보완 가능",
         "C · 재제출됨 (보라)",
         "없음",
         "재제출 항목 검토 후 승인 또는 추가 보완"],
        ["D", "개통 준비 완료",
         "items 전체 state === 'approved'",
         "1 (최고)",
         "✅ 개통 완료 처리 버튼 (초록 카드)\n취소는 작은 링크로 접근",
         "D · 개통 준비 완료 (초록)",
         "없음",
         "✅ 개통 완료 처리 → 케이스 종료"],
    ]

    status_bgs = {
        "null": GRAY_LT,
        "A": BLUE_LT,
        "B": AMBER_LIGHT,
        "C": PURPLE_LT,
        "D": EMERALD_LT,
    }

    for r_idx, row in enumerate(rows, 4):
        bg = status_bgs[row[0]]
        make_cell(ws, r_idx, 1, row[0], bg=bg, bold=True, align="center", size=12)
        for c_idx, val in enumerate(row[1:], 2):
            make_cell(ws, r_idx, c_idx, val, bg=bg)
        ws.row_dimensions[r_idx].height = 60

    ws.row_dimensions[9].height = 10
    note_text = ("💡 우선순위: D > B > C > A — Status는 직접 저장하지 않고 항목 상태에서 실시간 파생됨 "
                 "(src/lib/computeStatus.ts)")
    note = ws.cell(row=10, column=1, value=note_text)
    note.fill = fill(GRAY_LT)
    note.font = font(italic=True, color=GRAY, size=9)
    note.alignment = left()
    ws.merge_cells(start_row=10, start_column=1, end_row=10, end_column=len(headers))
    ws.row_dimensions[10].height = 22


# ══════════════════════════════════════════════════════════════════════════════
# 6. 자동 발송 메시지 (SCRIPT)
# ══════════════════════════════════════════════════════════════════════════════
def build_script(wb):
    ws = wb.create_sheet("6. 자동 발송 메시지")
    sheet_title(ws, "6. 자동 발송 메시지 (SCRIPT)", "오퍼레이터 자유 채팅 없음 — 트리거 기반 자동 발송")

    headers = ["#", "메시지명", "트리거 조건 (정확한 코드 조건)",
               "발신자 역할", "딜레이", "메시지 본문", "변수", "비고"]
    widths = [4, 18, 40, 12, 10, 50, 16, 24]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        make_header(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    rows = [
        ["1", "인트로\n(앱번호 경로)",
         "operatorChangeSeg() 호출\n& isFirstConfirmation === true\n(prev.segConfirmed === false)",
         "operator", "300ms",
         "안녕하세요, {name}님! 하이어비자 모바일 개통팀입니다.\n"
         "신청해 주셔서 감사합니다. 아래 서류 목록을 확인하시고\n"
         "순서대로 제출해 주세요. 궁금한 점이 있으시면 언제든지\n"
         "말씀해 주세요. 🙏",
         "{name} = studentName", "Seg 확인과 동시에 발송"],
        ["2", "인트로\n(워크플로우 경로)",
         "studentApplyAndStart(seg, simType) 호출\n(Q3 '신청 완료' 버튼 클릭)",
         "operator", "1800ms",
         "안녕하세요, {name}님! 하이어비자 모바일 개통팀입니다.\n"
         "신청해 주셔서 감사합니다. 아래 서류 목록을 확인하시고\n"
         "순서대로 제출해 주세요. 궁금한 점이 있으시면 언제든지\n"
         "말씀해 주세요. 🙏",
         "{name} = studentName", "긴 딜레이로 오퍼레이터 배정 시뮬레이션"],
        ["3", "재제출 확인",
         "studentSubmitItem() 완료\n& updatedItem.state === 'resubmitted'\n& newStatus === 'C'\n& prev.status !== 'C'",
         "operator", "600ms",
         "재제출해 주신 서류를 확인하겠습니다.\n잠시만 기다려 주세요. 🔍",
         "없음",
         "Status B→C 전환 시점 1회만 발송\n(여러 항목 재제출 시 마지막 항목 기준)"],
        ["4", "전체 승인 완료",
         "operatorApproveItem() 완료\n& newStatus === 'D'\n& prev.status !== 'D'",
         "operator", "400ms",
         "모든 서류 확인이 완료되었습니다!\n곧 개통 처리를 진행해 드리겠습니다. 📱",
         "없음",
         "Status D 전환 시점 1회만 발송"],
        ["5", "개통 완료",
         "operatorActivate() 호출\n(status === 'D' 조건 충족 시)",
         "operator", "즉시",
         "{name}님, 모바일 개통이 정상적으로 완료되었습니다! 🎉\n"
         "{segLabel} · {simLabel} 방식으로 개통 처리되었습니다.\n"
         "하이어비자 모바일 서비스 이용을 진심으로 환영합니다. 😊\n"
         "이용 중 불편하신 사항은 언제든지 채팅으로 문의해 주세요.",
         "{name}/{segLabel}/{simLabel}", "채팅 종료 처리 동시 실행"],
        ["6", "무응답 안내\n(대기 타이머)",
         "학생 메시지/제출 후 60초 경과\n(오퍼레이터 응답 없음)",
         "bot", "60초",
         "⏳ 곧 응대 예정입니다. 잠시만 기다려주세요.",
         "없음",
         "오퍼레이터 응답 시 타이머 리셋\n(메시지 전송 / 항목 처리 / Seg 확인)"],
        ["7", "보완 요청\n(항목별 템플릿)",
         "operatorRequestSupplement(code, category, detail) 호출\n(SuppModal에서 사유 선택 후 전송)",
         "operator (template)", "즉시",
         "[보완 요청] {항목명}\n{category} · {detail(선택)}",
         "{항목명}/{category}/{detail}", "학생 채팅에 amber 배경 카드로 표시"],
        ["8", "사증발급확인서 요청",
         "operatorAddVisaConfirm() 호출\n(체크리스트 우측 '🛂 사증발급확인서 요청' 버튼)",
         "operator (template)", "즉시",
         "[추가 서류 요청] 사증발급확인서\n자사에서 확보하지 못한 관계로\n사증발급확인서 제출을 요청드립니다.",
         "없음", "항목 state = supp_requested로 즉시 설정"],
        ["9", "상담사 연결 알림\n(시스템)",
         "studentQuickReply('상담사 연결이 필요해요', true) 호출\n(consultantMode = true 설정 동시)",
         "system (warn)", "즉시",
         "🔔 상담사 직접 연결이 요청되었습니다.\n담당 상담사가 곧 연결됩니다.",
         "없음", "오퍼레이터 자유 채팅창 활성화"],
    ]

    script_bgs = [
        BLUE_LT, NAVY_LIGHT, PURPLE_LT, EMERALD_LT, EMERALD_LT,
        ORANGE_LT, AMBER_LIGHT, AMBER_LIGHT, GRAY_LT
    ]

    for r_idx, (row, bg) in enumerate(zip(rows, script_bgs), 4):
        make_cell(ws, r_idx, 1, row[0], bg=bg, bold=True, align="center")
        make_cell(ws, r_idx, 2, row[1], bg=bg, bold=True)
        make_cell(ws, r_idx, 3, row[2], bg=GRAY_LT)    # condition (monospace feel)
        make_cell(ws, r_idx, 4, row[3], bg=bg, align="center")
        make_cell(ws, r_idx, 5, row[4], bg=bg, align="center", bold=True)
        make_cell(ws, r_idx, 6, row[5], bg=bg)
        make_cell(ws, r_idx, 7, row[6], bg=bg, fg=NAVY)
        make_cell(ws, r_idx, 8, row[7], bg=WHITE, fg=GRAY)
        ws.row_dimensions[r_idx].height = 72


# ══════════════════════════════════════════════════════════════════════════════
# 7. 퀵리플라이 버튼
# ══════════════════════════════════════════════════════════════════════════════
def build_quickreply(wb):
    ws = wb.create_sheet("7. 퀵리플라이 버튼")
    sheet_title(ws, "7. 학생 퀵리플라이 버튼", "Status B 진입 시 자동 노출")

    headers = ["#", "노출 조건", "버튼 텍스트", "버튼 스타일",
               "전송 메시지", "requestConsultant", "시스템 액션",
               "오퍼레이터에게 보이는 것"]
    widths = [4, 30, 30, 16, 30, 16, 30, 28]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        make_header(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    cond = "status === 'B'\n& !consultantMode\n& !closed"

    rows = [
        ["1", cond,
         "❓ 어떤 부분이 부족한 건지 모르겠어요",
         "주황 아웃라인",
         "어떤 부분이 부족한 건지 모르겠어요",
         "false",
         "없음\n(메시지만 전송)",
         "학생 발송 메시지로 표시됨\n오퍼레이터가 수동 답변"],
        ["2", cond,
         "❓ 어떻게 제출해야 하는지 모르겠어요",
         "주황 아웃라인",
         "어떻게 제출해야 하는지 모르겠어요",
         "false",
         "없음\n(메시지만 전송)",
         "학생 발송 메시지로 표시됨\n오퍼레이터가 수동 답변"],
        ["3", cond,
         "💬 상담사 연결이 필요해요",
         "네이비 아웃라인",
         "상담사 연결이 필요해요",
         "true",
         "consultantMode = true\n시스템 알림 메시지 발송 (즉시)\n대기 타이머 시작",
         "오퍼레이터 자유 채팅 입력창 활성화\n학생 화면: '💬 상담사 연결 중' 배지"],
    ]

    row_bgs = [AMBER_LIGHT, AMBER_LIGHT, NAVY_LIGHT]
    for r_idx, (row, bg) in enumerate(zip(rows, row_bgs), 4):
        make_cell(ws, r_idx, 1, row[0], bg=bg, bold=True, align="center")
        for c_idx, val in enumerate(row[1:], 2):
            make_cell(ws, r_idx, c_idx, val, bg=bg)
        ws.row_dimensions[r_idx].height = 64

    # ── 미정의 영역 표 ──
    ws.row_dimensions[7].height = 12
    make_header(ws, 8, 1, "⚠️ 향후 정의 필요: 퀵리플라이 응답 스크립트", bg="DC2626", size=11)
    ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=8)
    ws.row_dimensions[8].height = 24

    undef_headers = ["버튼", "학생 메시지", "오퍼레이터 표준 응답 스크립트 (미정의)", "비고"]
    undef_widths = [30, 30, 50, 20]
    for col, (h, w) in enumerate(zip(undef_headers, undef_widths), 1):
        make_header(ws, 9, col, h, bg=RED_LT, fg="991B1B")
        ws.column_dimensions[get_column_letter(col)].width = w

    undef_rows = [
        ["❓ 어떤 부분이 부족한 건지 모르겠어요",
         "어떤 부분이 부족한 건지 모르겠어요",
         "(미정의 — 항목별 표준 응답 문구 필요)", "항목별 다를 수 있음"],
        ["❓ 어떻게 제출해야 하는지 모르겠어요",
         "어떻게 제출해야 하는지 모르겠어요",
         "(미정의 — 제출 방법 가이드 스크립트 필요)", "항목별 다를 수 있음"],
    ]
    for r_idx, row in enumerate(undef_rows, 10):
        for c_idx, val in enumerate(row, 1):
            make_cell(ws, r_idx, c_idx, val,
                      bg=RED_LT if c_idx == 3 else WHITE, fg="991B1B" if c_idx == 3 else "374151")
        ws.row_dimensions[r_idx].height = 36


# ══════════════════════════════════════════════════════════════════════════════
# 8. 보완 사유 카테고리
# ══════════════════════════════════════════════════════════════════════════════
def build_supp_categories(wb):
    ws = wb.create_sheet("8. 보완 사유 카테고리")
    sheet_title(ws, "8. 보완 사유 카테고리 (SuppModal)", "항목별 오퍼레이터 선택 가능 사유 목록")

    headers = ["항목 코드", "항목명", "아이콘", "#", "보완 사유 텍스트", "적용 Seg/SimType", "비고"]
    widths = [18, 20, 8, 4, 50, 20, 20]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        make_header(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    cats = [
        ("passport_scan", "여권 양면 평판 스캔본", "📘", "전체", [
            "사진이 흐릿함 - 재스캔 필요",
            "한쪽 면만 보임 - 양면 모두 포함 필요",
            "정보 일부 가림 - 정면 재스캔",
            "만료일 임박/지남 - 재발급 후 제출",
        ]),
        ("device_info", "기기정보 캡처본", "📱", "S1+S2 / eSIM 전용", [
            "IMEI가 보이지 않음 - *#06# 화면 재캡처",
            "모델명 확인 불가 - 설정 > 일반 > 정보 화면 캡처",
            "eSIM 지원 여부 확인 불가 - 설정에서 eSIM 항목 포함하여 재캡처",
        ]),
        ("sim_photo", "유심 사진", "💳", "S1+S2 / uSIM 전용", [
            "사진이 흐릿함 - 재촬영 필요",
            "유심 번호가 보이지 않음 - 번호 부분 포함하여 재촬영",
            "다른 유심이 촬영됨 - 실제 사용할 유심 재촬영",
        ]),
        ("current_carrier", "기존 통신사 확인", "🏢", "S2 전용", [
            "문자 캡처본이 불명확함 - 재캡처 필요",
            "통신사 정보가 보이지 않음",
            "다른 번호의 캡처본으로 보임 - 본인 번호로 재확인",
        ]),
        ("current_number", "기존 전화번호", "📞", "S2 전용", [
            "번호 형식 오류 - 올바른 번호 재입력",
            "통신사 캡처본의 번호와 불일치 - 재확인 필요",
            "본인 명의 아님 - 확인 필요",
        ]),
        ("visa_confirm", "사증발급확인서", "🛂", "선택적 (오퍼레이터 추가)", [
            "파일이 불명확함 - 재업로드 필요",
            "만료된 서류 - 최신 서류 필요",
            "다른 서류가 첨부됨 - 사증발급확인서로 재제출",
        ]),
    ]

    item_bgs = {
        "passport_scan": BLUE_LT,
        "device_info": PURPLE_LT,
        "sim_photo": TEAL_LT,
        "current_carrier": AMBER_LIGHT,
        "current_number": ORANGE_LT,
        "visa_confirm": GRAY_LT,
    }

    r_idx = 4
    for code, name, icon, seg, reasons in cats:
        bg = item_bgs[code]
        first = True
        for i, reason in enumerate(reasons, 1):
            make_cell(ws, r_idx, 1, code if first else "", bg=bg, bold=first)
            make_cell(ws, r_idx, 2, name if first else "", bg=bg, bold=first)
            make_cell(ws, r_idx, 3, icon if first else "", bg=bg, align="center", size=14)
            make_cell(ws, r_idx, 4, str(i), bg=bg, align="center")
            make_cell(ws, r_idx, 5, reason, bg=bg)
            make_cell(ws, r_idx, 6, seg if first else "", bg=bg)
            make_cell(ws, r_idx, 7,
                      "기타(직접 입력) 옵션 항상 포함" if first else "",
                      bg=WHITE, fg=GRAY)
            ws.row_dimensions[r_idx].height = 28
            r_idx += 1
            first = False

    note_r = r_idx + 1
    note = ws.cell(row=note_r, column=1,
                   value="💡 SuppModal에서 '기타 (직접 입력)' 선택 시 오퍼레이터가 자유 텍스트 입력 가능. "
                         "카테고리는 config/segments.json에서 관리.")
    note.fill = fill(GRAY_LT)
    note.font = font(italic=True, color=GRAY, size=9)
    note.alignment = left()
    ws.merge_cells(start_row=note_r, start_column=1, end_row=note_r, end_column=7)
    ws.row_dimensions[note_r].height = 22


# ══════════════════════════════════════════════════════════════════════════════
# 9. 미정의 항목 (Backlog)
# ══════════════════════════════════════════════════════════════════════════════
def build_backlog(wb):
    ws = wb.create_sheet("9. 향후 정의 필요 항목")
    sheet_title(ws, "9. 향후 정의 필요 항목 (Product Backlog)",
                "현재 구현은 있으나 운영 기준 미확정")

    headers = ["#", "영역", "현재 상태 (코드)", "필요한 정의", "우선순위", "담당 제안"]
    widths = [4, 20, 36, 40, 12, 16]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        make_header(ws, 3, col, h)
        ws.column_dimensions[get_column_letter(col)].width = w

    rows = [
        ["1", "퀵리플라이 응답",
         "consultantMode 없이 버튼 클릭 시\n오퍼레이터가 자유 응답",
         "버튼별 표준 응답 스크립트 정의\n(❓ 어떤 부분 / ❓ 제출 방법 각각)",
         "🔴 높음", "PM + CS 팀"],
        ["2", "보완요청 추가 안내 멘트",
         "SuppModal → 카테고리 + 자유 텍스트\n→ 템플릿 메시지로 발송",
         "항목별 보완 시 표준 안내 문구\n(예: 양면 재스캔 방법 안내 등)",
         "🔴 높음", "PM + CS 팀"],
        ["3", "워크플로우 서류 제출 UX",
         "학생 채팅 하단 항목별 제출 버튼\n(ItemInputModal로 파일 업로드)",
         "항목별 업로드 전 가이드 메시지\n스텝별 안내 여부",
         "🟡 중간", "UX + 개발"],
        ["4", "보완 미재제출 에스컬레이션",
         "미구현 (타이머/알림 없음)",
         "재제출 요청 후 N시간 미응답 시\n자동 알림 또는 케이스 처리 기준",
         "🟡 중간", "PM + 개발"],
        ["5", "Seg 수동 변경 시 기존 승인 처리",
         "공통 항목 상태 유지\n신규 항목은 empty로 추가",
         "기존 승인 항목을 그대로 유지할지\n또는 재심사 시작할지 기준",
         "🟡 중간", "PM"],
        ["6", "사증발급확인서 확보 기준",
         "오퍼레이터 판단으로 직접 버튼 클릭",
         "자사 확보 가능한 비자 유형 목록\n(확보 시 학생에게 요청 안 함)",
         "🟢 낮음", "운영팀"],
        ["7", "개통 완료 후 안내",
         "임의 텍스트 하드코딩\n(SCRIPT.activated)",
         "실제 서비스 안내 URL/메시지 확정\n(요금제 안내, 고객센터 번호 등)",
         "🟢 낮음", "PM + 마케팅"],
        ["8", "신청번호 형식 검증",
         "대문자 변환만 처리\n(HV2026- 형식 검증 없음)",
         "허용 형식 및 오류 메시지 정의",
         "🟢 낮음", "개발"],
    ]

    prio_colors = {"🔴 높음": RED_LT, "🟡 중간": AMBER_LIGHT, "🟢 낮음": EMERALD_LT}
    for r_idx, row in enumerate(rows, 4):
        prio_bg = prio_colors.get(row[4], WHITE)
        make_cell(ws, r_idx, 1, row[0], bg=GRAY_LT, bold=True, align="center")
        make_cell(ws, r_idx, 2, row[1], bg=GRAY_LT, bold=True)
        make_cell(ws, r_idx, 3, row[2], bg=WHITE, fg=GRAY)
        make_cell(ws, r_idx, 4, row[3], bg=WHITE, bold=True)
        make_cell(ws, r_idx, 5, row[4], bg=prio_bg, align="center", bold=True)
        make_cell(ws, r_idx, 6, row[5], bg=WHITE, fg=GRAY)
        ws.row_dimensions[r_idx].height = 48


# ══════════════════════════════════════════════════════════════════════════════
# Main
# ══════════════════════════════════════════════════════════════════════════════
def main():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    build_entry_path(wb)
    build_workflow(wb)
    build_item_matrix(wb)
    build_item_states(wb)
    build_status(wb)
    build_script(wb)
    build_quickreply(wb)
    build_supp_categories(wb)
    build_backlog(wb)

    out = "/home/user/cs_mobile_chat/docs/workflow-branches.xlsx"
    wb.save(out)
    print(f"✅ Saved: {out}")
    print(f"   Sheets: {[ws.title for ws in wb.worksheets]}")


if __name__ == "__main__":
    main()
